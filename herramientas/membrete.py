"""Genera la hoja membretada en blanco y el modelo de presupuesto (.docx, A4)."""
import pathlib, sys
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

sys.path.insert(0, str(pathlib.Path(__file__).parent))
from datos import *  # noqa: E402,F403

SALIDA = pathlib.Path(__file__).resolve().parent.parent / "membrete"
TINTA_C, DORADO_C, GRIS = RGBColor.from_string(TINTA), RGBColor.from_string(DORADO), RGBColor(0x5B, 0x60, 0x6B)
SERIF, SANS = "Georgia", "Calibri"


def run(p, texto, tam=10, color=TINTA_C, fuente=SANS, negrita=False, cursiva=False, espaciado=0):
    r = p.add_run(texto)
    r.font.name, r.font.size, r.font.color.rgb = fuente, Pt(tam), color
    r.font.bold, r.font.italic = negrita, cursiva
    r._element.rPr.rFonts.set(qn("w:eastAsia"), fuente)
    if espaciado:
        sp = OxmlElement("w:spacing"); sp.set(qn("w:val"), str(espaciado)); r._element.rPr.append(sp)
    return r


def borde(p, lado="bottom", color=DORADO, grosor=6):
    pPr = p._p.get_or_add_pPr()
    bdr = OxmlElement("w:pBdr"); b = OxmlElement(f"w:{lado}")
    for k, v in (("val", "single"), ("sz", str(grosor)), ("space", "6"), ("color", color)):
        b.set(qn(f"w:{k}"), v)
    bdr.append(b); pPr.append(bdr)


def fmt(p, antes=0, despues=0, alin=None):
    p.paragraph_format.space_before, p.paragraph_format.space_after = Pt(antes), Pt(despues)
    if alin is not None:
        p.alignment = alin


def base():
    d = Document()
    s = d.sections[0]
    s.page_width, s.page_height = Cm(21), Cm(29.7)
    s.top_margin, s.bottom_margin, s.left_margin, s.right_margin = Cm(4.2), Cm(3), Cm(2.5), Cm(2.5)
    s.header_distance, s.footer_distance = Cm(1.4), Cm(1.2)
    st = d.styles["Normal"]; st.font.name, st.font.size = SANS, Pt(10.5); st.font.color.rgb = TINTA_C
    st.element.rPr.rFonts.set(qn("w:eastAsia"), SANS)

    # Encabezado: nombre a la izquierda, rubro a la derecha, filete dorado
    h = s.header; h.paragraphs[0].text = ""
    t = h.add_table(rows=1, cols=2, width=Cm(16)); t.alignment = WD_TABLE_ALIGNMENT.CENTER
    anchos(t, 10.3, 5.7)
    izq, der = t.rows[0].cells
    p = izq.paragraphs[0]; fmt(p)
    run(p, "Wasiluk Yasinsky ", 18, fuente=SERIF); run(p, "&", 18, DORADO_C, SERIF, cursiva=True); run(p, " Asociados", 18, fuente=SERIF)
    p = der.paragraphs[0]; fmt(p, antes=4, alin=WD_ALIGN_PARAGRAPH.RIGHT)
    run(p, RUBRO.upper(), 7, GRIS, espaciado=10)
    p = der.add_paragraph(); fmt(p, alin=WD_ALIGN_PARAGRAPH.RIGHT)
    run(p, CIUDAD, 8, GRIS)
    linea = h.add_paragraph(); fmt(linea, despues=0); borde(linea)
    h.paragraphs[0]._p.getparent().remove(h.paragraphs[0]._p)  # párrafo vacío inicial

    # Pie: personas y contacto
    f = s.footer
    p = f.paragraphs[0]; fmt(p, despues=3, alin=WD_ALIGN_PARAGRAPH.CENTER); borde(p, "top", grosor=4)
    for i, q in enumerate(PERSONAS):
        if i: run(p, "     ·     ", 8, DORADO_C)
        run(p, f"{q['titulo']} {q['nombre']} {q['apellido']}", 8, negrita=True)
        run(p, f"  {q['cargo']}  ·  {q['tel']}", 8, GRIS)
    p = f.add_paragraph(); fmt(p, alin=WD_ALIGN_PARAGRAPH.CENTER)
    run(p, CORREO, 8, GRIS); run(p, "   ·   ", 8, DORADO_C); run(p, WEB, 8, GRIS)
    if RUT:
        run(p, "   ·   ", 8, DORADO_C); run(p, f"RUT {RUT}", 8, GRIS)
    return d


def anchos(t, *cm):
    """Fija anchos de columna (Word y LibreOffice leen cosas distintas: se ponen las dos)."""
    t.autofit = False
    for gc, w in zip(t._tbl.tblGrid.findall(qn("w:gridCol")), cm):
        gc.set(qn("w:w"), str(int(w * 567)))
    for fila in t.rows:
        for c, w in zip(fila.cells, cm):
            c.width = Cm(w)


def parrafo(d, texto="", tam=10.5, antes=0, despues=6, **kw):
    p = d.add_paragraph(); fmt(p, antes, despues)
    if texto: run(p, texto, tam, **kw)
    return p


def sombrear(celda, color):
    tcPr = celda._tc.get_or_add_tcPr(); sh = OxmlElement("w:shd")
    sh.set(qn("w:val"), "clear"); sh.set(qn("w:color"), "auto"); sh.set(qn("w:fill"), color); tcPr.append(sh)


def main():
    # 1) Hoja en blanco
    d = base()
    parrafo(d, "Punta del Este, ___ de __________ de 20__", color=GRIS)
    d.save(SALIDA / "Hoja membretada.docx")

    # 2) Modelo de presupuesto
    d = base()
    p = parrafo(d, despues=2); run(p, "PRESUPUESTO", 9, DORADO_C, negrita=True, espaciado=60)
    p = parrafo(d, despues=14); run(p, "Nº 0000", 22, fuente=SERIF)
    p.add_run("\t"); run(p, "Punta del Este, ___ de __________ de 20__", 10, GRIS)
    p.paragraph_format.tab_stops.add_tab_stop(Cm(16), alignment=2)

    datos = [("Cliente", "[Nombre o razón social]"), ("RUT / CI", "[Número]"),
             ("Contacto", "[Teléfono · correo]"), ("Asunto", "[Trámite o servicio]")]
    t = d.add_table(rows=len(datos), cols=2)
    anchos(t, 3, 13)
    for (k, v), fila in zip(datos, t.rows):
        run(fila.cells[0].paragraphs[0], k, 9, GRIS)
        run(fila.cells[1].paragraphs[0], v, 10.5)

    parrafo(d, "De acuerdo a lo conversado, detallamos los honorarios por los servicios solicitados:", antes=16, despues=10)

    items = [("[Descripción del servicio]", ""), ("[Descripción del servicio]", ""), ("[Descripción del servicio]", "")]
    t = d.add_table(rows=1 + len(items) + 3, cols=2); t.alignment = WD_TABLE_ALIGNMENT.CENTER
    enc = t.rows[0].cells
    for c, txt, al in ((enc[0], "Concepto", None), (enc[1], "Importe", WD_ALIGN_PARAGRAPH.RIGHT)):
        sombrear(c, TINTA); pp = c.paragraphs[0]; fmt(pp, 3, 3, al); run(pp, txt, 9, RGBColor(0xF4, 0xF1, 0xEA), negrita=True)
    for i, (concepto, imp) in enumerate(items, 1):
        c0, c1 = t.rows[i].cells
        fmt(c0.paragraphs[0], 4, 4); run(c0.paragraphs[0], concepto, 10.5)
        fmt(c1.paragraphs[0], 4, 4, WD_ALIGN_PARAGRAPH.RIGHT); run(c1.paragraphs[0], "$ 0", 10.5)
        borde(c0.paragraphs[0], color="DDD8CC", grosor=4); borde(c1.paragraphs[0], color="DDD8CC", grosor=4)
    for j, (txt, fuerte) in enumerate((("Subtotal", False), ("IVA 22 %", False), ("Total", True))):
        c0, c1 = t.rows[1 + len(items) + j].cells
        fmt(c0.paragraphs[0], 4, 2, WD_ALIGN_PARAGRAPH.RIGHT); fmt(c1.paragraphs[0], 4, 2, WD_ALIGN_PARAGRAPH.RIGHT)
        run(c0.paragraphs[0], txt, 11 if fuerte else 10, negrita=fuerte, color=TINTA_C if fuerte else GRIS)
        run(c1.paragraphs[0], "$ 0", 12 if fuerte else 10.5, negrita=fuerte)
        if fuerte: sombrear(c0, "F4F1EA"); sombrear(c1, "F4F1EA")
    anchos(t, 12, 4)

    p = parrafo(d, antes=18, despues=4); run(p, "CONDICIONES", 9, DORADO_C, negrita=True, espaciado=60)
    for c in ("Validez de la propuesta: 15 días desde la fecha de emisión.",
              "Forma de pago: [50 % al inicio y 50 % al finalizar / contado / mensual].",
              "No incluye: tributos, timbres, tasas registrales ni gastos de terceros, que se liquidan aparte según comprobante.",
              "Moneda: [pesos uruguayos / dólares estadounidenses]."):
        p = d.add_paragraph(style="List Bullet"); fmt(p, despues=2); run(p, c, 10)

    parrafo(d, "Quedamos a su disposición para cualquier consulta.", antes=14, despues=36)
    p = parrafo(d, despues=0); run(p, "______________________________", 10, GRIS)
    p = parrafo(d, despues=0); run(p, "[Esc. Lorena Yasinsky / Cr. José Wasiluk]", 10.5, fuente=SERIF)
    p = parrafo(d); run(p, NOMBRE, 9, GRIS)
    d.save(SALIDA / "Presupuesto (modelo).docx")
    print("ok", SALIDA)


if __name__ == "__main__":
    main()
